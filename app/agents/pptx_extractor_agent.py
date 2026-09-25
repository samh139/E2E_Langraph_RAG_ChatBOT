import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.workflow.state import IngestionState
from app.llm.llm_client import OllamaClient
from app.app_logger import LoggerFactory
from app.storage.storage_service import StorageService # Import your MongoDB StorageService

class PPTXExtractorAgent:

    def __init__(self):
        # Decoupled state from initialization to match your updated orchestration schema
        self.logger = LoggerFactory.get_logger(__name__)
        self.ollama_cloud = OllamaClient()

    # MATCHES LANGGRAPH SIGNATURE: Accepts state and returns a state patch dictionary
    def execute(self, state: IngestionState) -> dict:
        # FIXED: Read values directly from Pydantic using dot-notation attributes
        file_path = state.file_path
        run_id = state.run_id
        self.logger.info(f"[PPTX Extractor] Processing: {file_path} | Run ID: {run_id}")

        # BLOCK 1: Open presentation and initialize tracking lists
        presentation = Presentation(file_path)
        slides = []       
        elements = []     
        all_text = []     

        image_counter = 0
        MIN_IMAGE_SIZE_BYTES = 1000 # Ignore layout graphics, slide template borders or lines

        # BLOCK 2: Step through slides (1-indexed for logging clarity)
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text = []

            # BLOCK 3: Step through every visual shape element present on the slide canvas
            for shape_number, shape in enumerate(slide.shapes, start=1):
                
                # --- NEW INLINE BRANCH: IMAGE CHECKING ---
                # Check if this shape is an explicit picture type or layout snapshot container
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_bytes = shape.image.blob

                    # Filter out minor graphical styling blocks
                    if len(image_bytes) < MIN_IMAGE_SIZE_BYTES:
                        continue

                    image_counter += 1

                    # Collect context from text extracted down the current slide so far
                    surrounding_context = "\n".join(slide_text).strip()

                    # Exact preserved prompt layout structure
                    prompt = (
                        "You are an expert document analysis assistant. Transcribe this embedded document image or diagram precisely.\n\n"
                        f"DOCUMENT CONTEXT LOCATION HINT:\n"
                        f"Section: Slide {slide_number}\n"
                        f"Sub-Section: Shape {shape_number}\n\n"
                        f"The text right above this image discusses:\n"
                        f"--- BEGIN CONTEXT ---\n"
                        f"{surrounding_context if surrounding_context else '[No slide text content available above this image]'}\n"
                        f"--- END CONTEXT ---\n\n"
                        "INSTRUCTIONS:\n"
                        "Using the layout location context above, analyze the chart, logic diagram, or code screenshot. "
                        "Precisely transcribe all syntax or structural code text into Markdown notation."
                    )

                    try:
                        self.logger.info(f"[PPTX Extractor] 🚀 Slide {slide_number} Image #{image_counter} Detected. Running vision inference...")
                        visual_description = self.ollama_cloud.analyze_vision(
                            image_bytes=image_bytes,
                            user_prompt=prompt,
                        )

                        if visual_description and visual_description.strip():
                            visual_element_text = f"\n[IMAGE ANALYSIS: {visual_description.strip()}]\n"
                            slide_text.append(visual_element_text)
                            elements.append({
                                "text": visual_element_text,
                                "element_type": "paragraph",
                                "page": slide_number,
                                "slide": slide_number,
                                "shape_index": shape_number,
                            })
                    except Exception as e:
                        self.logger.error(f"[PPTX Extractor Error] Visual parsing break on slide {slide_number}: {e}")

                    continue # Finished processing picture shape, jump to next element

                # --- ORIGINAL TEXT EXTRACTION LOGIC STAYS INTACT ---
                if not hasattr(shape, "text"):
                    continue

                text = (shape.text or "").strip()
                if not text:
                    continue

                slide_text.append(text)

                element_type = "paragraph"
                heading_level = None
                is_title = False

                # BLOCK 4: Title parsing
                try:
                    is_title = (
                        shape.is_placeholder
                        and shape.placeholder_format.type in (1, 3) 
                    )
                except Exception:
                    is_title = False 

                if is_title:
                    element_type = "heading"
                    heading_level = 1

                # BLOCK 5: Font sizing calculation logic
                font_sizes = []
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is not None:
                                font_sizes.append(run.font.size.pt) 

                max_font_size = max(font_sizes) if font_sizes else None

                # BLOCK 6: Append text element
                elements.append({
                    "text": text,
                    "element_type": element_type,
                    "heading_level": heading_level,
                    "page": slide_number,  
                    "slide": slide_number,
                    "shape_index": shape_number,
                    "font_size": max_font_size,
                    "is_title": is_title,
                })

            # BLOCK 7: Compile individual slide text data sets into parent matrices
            slide_content = "\n".join(slide_text)
            
            slides.append({
                "slide": slide_number,
                "text": slide_content,
            })

            if slide_content:
                all_text.append(slide_content)

        # BLOCK 8: Create a single unified master string and format output for LangGraph pipeline use
        extracted_text = "\n\n".join(all_text)

        self.logger.info(
            f"[PPTX Extractor] Processing Complete. "
            f"Slides={len(slides)} | "
            f"Images Extracted={image_counter} | "
            f"Characters={len(extracted_text)}"
        )

        extraction_metadata = {
            "slide_count": len(slides),
            "slides": slides,
            "elements": elements,
        }

        # Safe local DB saving mechanism
        try:
            # Construct dictionary payload manually to mirror your other extractors
            mongo_payload = {
                "run_id": run_id,
                "file_path": file_path,
                "file_name": state.file_name,
                "file_extension": state.file_extension,
                "mime_type": state.mime_type,
                "file_type": state.file_type,
                "extracted_text": extracted_text,
                "extraction_metadata": extraction_metadata,
                "status": "extracted",
                "error": None
            }
            mongo_id = StorageService.save_pptx_data(mongo_payload, run_id=run_id)
            self.logger.info(f"[PPTX Extractor] Successfully persisted data to MongoDB. Doc Reference ID: {mongo_id}")
        except Exception as storage_err:
            self.logger.error(f"[PPTX Extractor Storage Error] Failed to write out to local database: {storage_err}")

        # CORRECT: Returns a clean patch dictionary for LangGraph to merge into state automatically
        return {
            "extracted_text": extracted_text,
            "extraction_metadata": extraction_metadata,
            "status": "extracted",
            "error": None,
        }
